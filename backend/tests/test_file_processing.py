import io
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from app.core.config import settings
from app.core.error_codes import ErrorCode
from app.services.document_parser import ParserFactory, find_doc_converter
from app.services.file_storage import (
    DocumentConverterUnavailable,
    DocumentParseError,
    EmptyFileError,
    FileStorage,
    FileTooLargeError,
    InvalidFileMimeTypeError,
    UnsupportedFileTypeError,
    UploadValidationError,
)


class FileStorageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name) / "uploads"
        self.storage = FileStorage(self.root, max_size_bytes=8)

    def tearDown(self) -> None:
        self.directory.cleanup()

    def test_save_uses_uuid_name_and_strips_posix_and_windows_path_parts(self) -> None:
        stored = self.storage.save(
            io.BytesIO(b"law text"),
            r"C:\fakepath\..\labor-law.TXT",
            "text/plain; charset=utf-8",
        )

        self.assertEqual(stored.file_name, "labor-law.TXT")
        self.assertEqual(stored.file_type, "txt")
        self.assertRegex(stored.path.name, r"^[0-9a-f]{32}\.txt$")
        self.assertEqual(stored.path.parent, self.root.resolve())
        self.assertEqual(stored.path.read_bytes(), b"law text")
        self.assertEqual(stored.size_bytes, 8)
        self.assertEqual(list(self.root.glob(".upload-*.tmp")), [])

    def test_rejects_unsupported_extension_and_mime_mismatch(self) -> None:
        with self.assertRaises(UnsupportedFileTypeError) as unsupported:
            self.storage.save(io.BytesIO(b"data"), "law.rtf", "application/rtf")
        self.assertIsInstance(unsupported.exception, UploadValidationError)
        self.assertEqual(unsupported.exception.code, ErrorCode.UNSUPPORTED_FILE_TYPE)

        with self.assertRaises(InvalidFileMimeTypeError) as mismatch:
            self.storage.save(io.BytesIO(b"data"), "law.pdf", "text/plain")
        self.assertEqual(mismatch.exception.http_status, 400)
        self.assertEqual(mismatch.exception.code, ErrorCode.INVALID_REQUEST)
        self.assertFalse(self.root.exists())

    def test_rejects_oversized_file_and_cleans_temporary_file(self) -> None:
        with self.assertRaises(FileTooLargeError) as too_large:
            self.storage.save(io.BytesIO(b"123456789"), "law.txt", "text/plain")

        self.assertEqual(too_large.exception.http_status, 413)
        self.assertEqual(too_large.exception.code, ErrorCode.FILE_TOO_LARGE)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_rejects_empty_file_without_leaving_a_file(self) -> None:
        with self.assertRaises(EmptyFileError) as empty:
            self.storage.save(io.BytesIO(b""), "law.txt", "text/plain")

        self.assertEqual(empty.exception.code, ErrorCode.EMPTY_FILE)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_failed_atomic_replace_does_not_leave_partial_files(self) -> None:
        with (
            patch("app.services.file_storage.os.replace", side_effect=OSError),
            self.assertRaises(OSError),
        ):
            self.storage.save(io.BytesIO(b"law"), "law.txt", "text/plain")

        self.assertEqual(list(self.root.iterdir()), [])

    def test_delete_only_removes_files_inside_upload_directory(self) -> None:
        stored = self.storage.save(io.BytesIO(b"law"), "law.txt", "text/plain")
        self.storage.delete(stored.path)
        self.assertFalse(stored.path.exists())

        outside = Path(self.directory.name) / "outside.txt"
        outside.write_text("keep", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.storage.delete(outside)
        self.assertTrue(outside.exists())


class ParserFactoryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)

    def tearDown(self) -> None:
        self.directory.cleanup()

    def test_txt_is_decoded_as_utf8_and_empty_text_is_rejected(self) -> None:
        source = self.root / "law.txt"
        source.write_bytes(b"\xef\xbb\xbfFirst\nSecond")
        self.assertEqual(ParserFactory.parse(source, "txt"), "First\nSecond")

        source.write_text("  \n", encoding="utf-8")
        with self.assertRaises(EmptyFileError):
            ParserFactory.parse(source, "txt")

    def test_txt_invalid_utf8_is_a_parse_error(self) -> None:
        source = self.root / "law.txt"
        source.write_bytes(b"\xff")
        with self.assertRaises(DocumentParseError):
            ParserFactory.parse(source, "txt")

    def test_pdf_pages_are_extracted_in_order(self) -> None:
        source = self.root / "law.pdf"
        source.write_bytes(b"pdf")
        pypdf = types.ModuleType("pypdf")
        pypdf.PdfReader = lambda _path: types.SimpleNamespace(
            pages=[
                types.SimpleNamespace(extract_text=lambda: "page one"),
                types.SimpleNamespace(extract_text=lambda: "page two"),
            ]
        )

        with patch.dict(sys.modules, {"pypdf": pypdf}):
            self.assertEqual(ParserFactory.parse(source, "pdf"), "page one\npage two")

    def test_pdf_parser_reads_a_real_pdf(self) -> None:
        from pypdf import PdfWriter
        from pypdf.generic import (
            DecodedStreamObject,
            DictionaryObject,
            NameObject,
        )

        source = self.root / "real-law.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (Page one) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        with source.open("wb") as output:
            writer.write(output)

        self.assertIn("Page one", ParserFactory.parse(source, "pdf"))

    def _fake_docx_modules(self) -> dict[str, types.ModuleType]:
        class Paragraph:
            def __init__(self, text: str) -> None:
                self.text = text

        class Cell:
            def __init__(self, text: str) -> None:
                self.text = text

        class Row:
            def __init__(self) -> None:
                self.cells = [Cell("table A"), Cell("table B")]

        class Table:
            def __init__(self) -> None:
                self.rows = [Row()]

        content = [Paragraph("before"), Table(), Paragraph("after")]
        document_module = types.ModuleType("docx")
        document_module.Document = lambda _path: types.SimpleNamespace(
            iter_inner_content=lambda: iter(content)
        )
        table_module = types.ModuleType("docx.table")
        table_module.Table = Table
        paragraph_module = types.ModuleType("docx.text.paragraph")
        paragraph_module.Paragraph = Paragraph
        text_module = types.ModuleType("docx.text")
        return {
            "docx": document_module,
            "docx.table": table_module,
            "docx.text": text_module,
            "docx.text.paragraph": paragraph_module,
        }

    def test_docx_paragraphs_and_tables_keep_document_order(self) -> None:
        source = self.root / "law.docx"
        source.write_bytes(b"docx")
        with patch.dict(sys.modules, self._fake_docx_modules()):
            parsed = ParserFactory.parse(source, "docx")
        self.assertEqual(parsed, "before\ntable A\ttable B\nafter")

    def test_docx_parser_reads_a_real_package(self) -> None:
        from docx import Document

        source = self.root / "real-law.docx"
        document = Document()
        document.add_paragraph("第一条 工资应当按时支付。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "劳动者"
        table.cell(0, 1).text = "依法享有休息权"
        document.add_paragraph("第二条 依法订立劳动合同。")
        document.save(source)

        self.assertEqual(
            ParserFactory.parse(source, "docx"),
            "第一条 工资应当按时支付。\n劳动者\t依法享有休息权\n第二条 依法订立劳动合同。",
        )

    def test_doc_conversion_uses_and_cleans_a_temporary_docx(self) -> None:
        source = self.root / "law.doc"
        source.write_bytes(b"doc")
        modules = self._fake_docx_modules()
        conversion_directories: list[Path] = []

        def convert(
            command: list[str], **_kwargs: object
        ) -> subprocess.CompletedProcess[str]:
            output_directory = Path(command[command.index("--outdir") + 1])
            conversion_directories.append(output_directory)
            (output_directory / "law.docx").write_bytes(b"converted")
            return subprocess.CompletedProcess(command, 0, "", "")

        with (
            patch.dict(sys.modules, modules),
            patch.object(settings, "doc_parser_backend", "libreoffice"),
            patch(
                "app.services.document_parser.shutil.which",
                return_value="/usr/bin/soffice",
            ),
            patch("app.services.document_parser.subprocess.run", side_effect=convert),
        ):
            parsed = ParserFactory.parse(source, "doc")

        self.assertEqual(parsed, "before\ntable A\ttable B\nafter")
        self.assertEqual(len(conversion_directories), 1)
        self.assertFalse(conversion_directories[0].exists())

    def test_doc_conversion_uses_msdoc2docx_by_default_and_cleans_output(self) -> None:
        source = self.root / "law.doc"
        source.write_bytes(b"doc")
        modules = self._fake_docx_modules()
        output_directories: list[Path] = []

        def convert(_source: str | Path, destination: str | Path) -> types.SimpleNamespace:
            destination_path = Path(destination)
            output_directories.append(destination_path.parent)
            destination_path.write_bytes(b"converted")
            return types.SimpleNamespace(
                report=types.SimpleNamespace(diagnostics=[]),
            )

        doc2docx = types.ModuleType("doc2docx")
        doc2docx.convert = convert  # type: ignore[attr-defined]
        with (
            patch.dict(
                sys.modules,
                {**modules, "doc2docx": doc2docx},
            ),
            patch.object(settings, "doc_parser_backend", "package"),
        ):
            parsed = ParserFactory.parse(source, "doc")

        self.assertEqual(parsed, "before\ntable A\ttable B\nafter")
        self.assertEqual(len(output_directories), 1)
        self.assertFalse(output_directories[0].exists())

    def test_doc_conversion_rejects_error_diagnostics_from_msdoc2docx(self) -> None:
        source = self.root / "law.doc"
        source.write_bytes(b"doc")

        def convert(_source: str | Path, destination: str | Path) -> types.SimpleNamespace:
            Path(destination).write_bytes(b"converted")
            diagnostic = types.SimpleNamespace(severity="error")
            return types.SimpleNamespace(
                report=types.SimpleNamespace(diagnostics=[diagnostic]),
            )

        doc2docx = types.ModuleType("doc2docx")
        doc2docx.convert = convert  # type: ignore[attr-defined]
        with (
            patch.dict(sys.modules, {"doc2docx": doc2docx}),
            patch.object(settings, "doc_parser_backend", "package"),
            self.assertRaisesRegex(DocumentParseError, "转换报告包含错误"),
        ):
            ParserFactory.parse(source, "doc")

    def test_doc_converter_handles_a_real_legacy_file_when_installed(self) -> None:
        from docx import Document

        converter = find_doc_converter()
        if converter is None:
            self.skipTest("LibreOffice is not installed")

        source = self.root / "law.docx"
        document = Document()
        document.add_paragraph("第一条 工资应当按时支付。")
        document.save(source)
        output_directory = self.root / "legacy"
        output_directory.mkdir()
        profile = self.root / "lo-profile"
        profile.mkdir()
        subprocess.run(
            [
                converter,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "doc",
                "--outdir",
                str(output_directory),
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=60,
        )

        legacy_source = output_directory / "law.doc"
        self.assertTrue(legacy_source.is_file())
        self.assertIn(
            "第一条 工资应当按时支付。", ParserFactory.parse(legacy_source, "doc")
        )

    def test_missing_doc_parser_and_converter_have_a_distinct_exception(self) -> None:
        source = self.root / "law.doc"
        source.write_bytes(b"doc")
        with (
            patch.object(settings, "doc_parser_backend", "package"),
            patch(
                "app.services.document_parser.package_doc_parser_available",
                return_value=False,
            ),
            patch("app.services.document_parser.shutil.which", return_value=None),
            self.assertRaises(DocumentConverterUnavailable),
        ):
            ParserFactory.parse(source, "doc")


if __name__ == "__main__":
    unittest.main()
